"""
Lógica de negocio: lectura de la hoja OFICIAL, parseo de lotes, cálculo del
tenor del nuevo lote y generación del documento (.xlsx, importable en Google Sheets).
Sin dependencias de interfaz, para poder probarse por separado.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HOJA = "OFICIAL"
COL_LOTE = "LOTE"
COL_TENOR = "Tenor Au Quintana (g/t)"


class ErrorDatos(Exception):
    """Error de validación con mensaje claro para el usuario."""


# --------------------------------------------------------------------------- #
# Utilidades de normalización
# --------------------------------------------------------------------------- #
def _norm(texto) -> str:
    """Minúsculas, sin tildes y sin espacios repetidos (para comparar encabezados)."""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().lower()


def a_numero(valor) -> float | None:
    """
    Convierte texto/número a float manejando separadores decimales.
    Acepta: 12.5 | "12,5" | "1.234,56" | "1,234.56" | "1 234,5".
    Devuelve None si está vacío o no es numérico.
    """
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    s = str(valor).strip().replace("\u00a0", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        # El último separador que aparece es el decimal.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def clave_lote(valor) -> str:
    """Clave de comparación de un LOTE: 601, 601.0 y '601' son el mismo lote."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip().upper()


# --------------------------------------------------------------------------- #
# Lectura del Excel
# --------------------------------------------------------------------------- #
def leer_hoja_oficial(archivo) -> pd.DataFrame:
    """
    Lee SOLO la hoja OFICIAL. `archivo` es una ruta o un objeto tipo archivo.
    Devuelve un DataFrame con columnas: LOTE, TENOR (texto original),
    FILA_EXCEL (número de fila real en la hoja).
    Detecta la fila de encabezados (primeras 30 filas) por si no es la fila 1.
    """
    try:
        xls = pd.ExcelFile(archivo, engine="openpyxl")
    except FileNotFoundError:
        raise ErrorDatos("No se encontró el archivo Excel indicado.")
    except Exception as e:
        raise ErrorDatos(f"No se pudo abrir el archivo como Excel (.xlsx): {e}")

    hoja = next((h for h in xls.sheet_names if h.strip().upper() == HOJA), None)
    if hoja is None:
        raise ErrorDatos(
            f"El archivo no tiene una hoja llamada '{HOJA}'. "
            f"Hojas encontradas: {', '.join(xls.sheet_names)}."
        )

    crudo = pd.read_excel(xls, sheet_name=hoja, header=None, dtype=object)
    n_lote, n_tenor = _norm(COL_LOTE), _norm(COL_TENOR)

    fila_enc = None
    for i in range(min(30, len(crudo))):
        celdas = [_norm(c) for c in crudo.iloc[i].tolist() if not pd.isna(c)]
        if n_lote in celdas and n_tenor in celdas:
            fila_enc = i
            break
    if fila_enc is None:
        # Diagnóstico: ¿qué columnas faltan en la fila 1?
        enc = [_norm(c) for c in crudo.iloc[0].tolist()] if len(crudo) else []
        faltan = [c for c, n in ((COL_LOTE, n_lote), (COL_TENOR, n_tenor)) if n not in enc]
        raise ErrorDatos(
            f"En la hoja '{HOJA}' faltan las columnas requeridas: {', '.join(faltan) or 'ambas'}. "
            f"Se esperan exactamente '{COL_LOTE}' y '{COL_TENOR}'."
        )

    encabezados = [_norm(c) if not pd.isna(c) else "" for c in crudo.iloc[fila_enc].tolist()]
    i_lote, i_tenor = encabezados.index(n_lote), encabezados.index(n_tenor)

    datos = crudo.iloc[fila_enc + 1 :].copy()
    df = pd.DataFrame(
        {
            "LOTE": datos.iloc[:, i_lote].values,
            "TENOR": datos.iloc[:, i_tenor].values,
            "FILA_EXCEL": datos.index + 1,  # índice 0-based -> fila de Excel 1-based
        }
    )
    df = df[df["LOTE"].notna() & (df["LOTE"].astype(str).str.strip() != "")].reset_index(drop=True)
    if df.empty:
        raise ErrorDatos(f"La hoja '{HOJA}' no contiene lotes debajo de los encabezados.")
    return df


# --------------------------------------------------------------------------- #
# Parseo de la entrada de lotes
# --------------------------------------------------------------------------- #
def parsear_seleccion(texto: str) -> list[str]:
    """
    Convierte "601-607" o "601, 605, 610" (o combinaciones "601-603, 610")
    en una lista de claves de lote (texto). Los rangos se expanden de a 1.
    Mantiene el orden y elimina repetidos.
    """
    if texto is None or not texto.strip():
        raise ErrorDatos("Ingresa los lotes que componen el nuevo lote (ej. 601-607 o 601, 605, 610).")

    tokens = [t.strip() for t in re.split(r"[,;\n]+", texto) if t.strip()]
    resultado: list[str] = []
    for tok in tokens:
        m = re.fullmatch(r"(\d+)\s*(?:-|–|—|a)\s*(\d+)", tok, flags=re.IGNORECASE)
        if m:
            ini, fin = int(m.group(1)), int(m.group(2))
            if ini > fin:
                raise ErrorDatos(f"Rango inválido '{tok}': el inicio es mayor que el final.")
            if fin - ini > 5000:
                raise ErrorDatos(f"El rango '{tok}' es demasiado grande (máx. 5001 lotes).")
            resultado.extend(str(n) for n in range(ini, fin + 1))
        elif re.fullmatch(r"[\w\-\.\/]+", tok):
            resultado.append(clave_lote(tok))
        else:
            raise ErrorDatos(f"No se entiende '{tok}'. Usa rangos (601-607) o listas (601, 605, 610).")

    vistos, unicos = set(), []
    for r in resultado:
        if r not in vistos:
            vistos.add(r)
            unicos.append(r)
    return unicos


def validar_peso(texto: str) -> float:
    """Valida que el peso sea un número positivo (acepta coma o punto decimal)."""
    p = a_numero(texto)
    if p is None:
        raise ErrorDatos("El peso aproximado debe ser un número (ej. 1500 o 1500,5).")
    if p <= 0:
        raise ErrorDatos("El peso aproximado debe ser un número mayor que cero.")
    return p


# --------------------------------------------------------------------------- #
# Cálculo
# --------------------------------------------------------------------------- #
@dataclass
class ResGrupo:
    """Resultado de un rango/grupo de lotes con su propio peso."""
    nombre: str                      # "Rango 1", "Rango 2", ...
    texto: str                       # texto ingresado por el usuario (ej. "601-612")
    lotes: list[str]
    tenores: list[float]             # tenor (g/t) de cada lote, mismo orden
    cantidad: int                    # n del grupo
    peso: float                      # peso asignado a este grupo
    peso_por_lote: float             # peso / n
    aportes: list[float]             # peso_por_lote * tenor_i
    tenor_grupo: float               # Σ aportes / peso (tenor del grupo solo)


@dataclass
class Resultado:
    grupos: list[ResGrupo]
    lotes: list[str]                 # todos los lotes, en orden de grupo
    tenores: list[float]
    cantidad: int                    # total de lotes tomados
    peso_total: float                # Σ pesos de los grupos
    suma_aportes: float              # Σ de todos los aportes
    tenor_nuevo: float               # suma_aportes / peso_total (g/t)
    avisos: list[str] = field(default_factory=list)


def calcular_tenor(tenores: list[float], peso: float) -> tuple[float, float, list[float], float]:
    """
    Cálculo para UN grupo de lotes (misma fórmula especificada):
      1) peso_por_lote = peso / n
      2) aporte_i      = peso_por_lote * tenor_i
      3) suma_aportes  = Σ aporte_i
      4) tenor         = suma_aportes / peso
    Devuelve (tenor, peso_por_lote, aportes, suma_aportes).
    """
    n = len(tenores)
    peso_por_lote = peso / n
    aportes = [peso_por_lote * t for t in tenores]
    suma_aportes = sum(aportes)
    return suma_aportes / peso, peso_por_lote, aportes, suma_aportes


def calcular_grupos(grupos: list[tuple[list[float], float]]) -> float:
    """
    Tenor del nuevo lote cuando hay varios grupos, cada uno con su propio peso:
      tenor_nuevo = Σ_grupos Σ_lotes (peso_grupo / n_grupo × tenor_i) / Σ peso_grupo
    Con un solo grupo equivale a `calcular_tenor`.
    """
    suma_aportes = sum(calcular_tenor(t, p)[3] for t, p in grupos)
    return suma_aportes / sum(p for _, p in grupos)


def resolver_lotes(df: pd.DataFrame, seleccion: list[str], modo: str) -> pd.DataFrame:
    """
    Devuelve las filas del DataFrame correspondientes a la selección.
      modo = "lote" -> los números se comparan con el valor de la columna LOTE.
      modo = "fila" -> los números son filas de la hoja de Excel (601-607 = filas 601 a 607).
    Lanza ErrorDatos si algún lote/fila no existe, o si hay tenores vacíos/no numéricos.
    """
    if modo == "fila":
        filas = {}
        for s in seleccion:
            if not s.isdigit():
                raise ErrorDatos(f"En modo 'fila de Excel' solo se aceptan números; recibí '{s}'.")
            filas[int(s)] = s
        por_fila = df.set_index("FILA_EXCEL")
        faltan = [str(f) for f in filas if f not in por_fila.index]
        if faltan:
            raise ErrorDatos(
                f"Las siguientes filas no existen (o no tienen lote) en la hoja {HOJA}: {', '.join(faltan)}."
            )
        sel = por_fila.loc[list(filas)].reset_index()
    else:
        claves = df["LOTE"].map(clave_lote)
        indice = {}
        for k, i in zip(claves, df.index):
            indice.setdefault(k, i)  # ante duplicados, se usa la primera aparición
        faltan = [s for s in seleccion if s not in indice]
        if faltan:
            raise ErrorDatos(
                f"Los siguientes lotes no existen en la hoja {HOJA}: {', '.join(faltan)}."
            )
        sel = df.loc[[indice[s] for s in seleccion]].reset_index(drop=True)

    sel = sel.copy()
    sel["TENOR_NUM"] = sel["TENOR"].map(a_numero)
    malos = sel[sel["TENOR_NUM"].isna()]
    if not malos.empty:
        det = ", ".join(f"{clave_lote(r.LOTE)} (fila {r.FILA_EXCEL}, valor: {r.TENOR!r})" for r in malos.itertuples())
        raise ErrorDatos(f"Tenor vacío o no numérico en: {det}.")
    return sel


def procesar(df: pd.DataFrame, grupos_in: list[tuple[str, str]], modo: str = "lote") -> Resultado:
    """
    Flujo completo. `grupos_in` es una lista de (texto_lotes, texto_peso):
      - un solo elemento  -> el peso se reparte en partes iguales entre sus lotes (caso simple).
      - varios elementos  -> cada rango tiene su propio peso; el tenor final es el promedio
                             ponderado por peso de todos los lotes.
    Un lote no puede estar en dos rangos a la vez (se contaría doble).
    """
    if not grupos_in:
        raise ErrorDatos("Ingresa al menos un rango de lotes con su peso.")
    multiple = len(grupos_in) > 1

    grupos: list[ResGrupo] = []
    for k, (txt_lotes, txt_peso) in enumerate(grupos_in, start=1):
        prefijo = f"Rango {k}: " if multiple else ""
        try:
            seleccion = parsear_seleccion(txt_lotes)
            peso = validar_peso(txt_peso)
            sel = resolver_lotes(df, seleccion, modo)
        except ErrorDatos as e:
            raise ErrorDatos(prefijo + str(e)) from None
        lotes = [clave_lote(v) for v in sel["LOTE"]]
        tenores = [float(t) for t in sel["TENOR_NUM"]]
        tenor_g, ppl, aportes, _ = calcular_tenor(tenores, peso)
        grupos.append(ResGrupo(f"Rango {k}", txt_lotes.strip(), lotes, tenores,
                               len(lotes), peso, ppl, aportes, tenor_g))

    todos = [l for g in grupos for l in g.lotes]
    repetidos = sorted({l for l in todos if todos.count(l) > 1})
    if repetidos:
        raise ErrorDatos(
            "Los siguientes lotes están en más de un rango y se contarían doble: "
            + ", ".join(repetidos) + ". Quítalos de uno de los rangos."
        )

    peso_total = sum(g.peso for g in grupos)
    suma = sum(a for g in grupos for a in g.aportes)
    avisos = []
    claves_hoja = df["LOTE"].map(clave_lote)
    if modo == "lote" and (claves_hoja.duplicated(keep=False) & claves_hoja.isin(todos)).any():
        avisos.append("Algún lote seleccionado aparece más de una vez en la hoja; se usó su primera aparición.")
    return Resultado(grupos, todos, [t for g in grupos for t in g.tenores], len(todos),
                     peso_total, suma, suma / peso_total, avisos)


# --------------------------------------------------------------------------- #
# Documento de salida (.xlsx -> abre directo en Google Sheets)
# --------------------------------------------------------------------------- #
def generar_documento(r: Resultado, id_nuevo_lote: str, unidad_peso: str = "") -> bytes:
    """
    Libro con el resultado. Las celdas calculadas son FÓRMULAS (Google Sheets/Excel
    las recalculan al abrir). Estructura:
      - Resumen (id, fecha, peso total, lotes tomados, tenor)
      - Tabla de rangos (peso de cada rango y cantidad de lotes)
      - Tabla de lotes (tenor, peso por lote y aporte)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Nuevo lote"
    f = "Arial"
    azul = Font(name=f, color="0000FF")
    normal, negrita = Font(name=f), Font(name=f, bold=True)
    blanca = Font(name=f, bold=True, color="FFFFFF")
    cab = PatternFill("solid", fgColor="1F3864")
    fino = Side(style="thin", color="BFBFBF")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    g = len(r.grupos)

    # Posiciones de las tablas
    g_ini = 11                       # primera fila de datos de la tabla de rangos
    g_fin = g_ini + g - 1
    l_cab = g_fin + 3                # encabezado de la tabla de lotes
    l_ini = l_cab + 1
    l_fin = l_ini + r.cantidad - 1
    u = f" ({unidad_peso})" if unidad_peso else ""

    ws["A1"] = "Resultado – Nuevo lote compuesto"
    ws["A1"].font = Font(name=f, bold=True, size=14)

    resumen = [
        ("Identificador del nuevo lote", id_nuevo_lote, azul, None),
        ("Fecha de generación", datetime.now().strftime("%Y-%m-%d %H:%M"), normal, None),
        ("Peso aproximado total" + u, f"=SUM(C{g_ini}:C{g_fin})", normal, "#,##0.00"),
        ("Cantidad de lotes tomados", f"=COUNTA(B{l_ini}:B{l_fin})", normal, "0"),
        ("Tenor del nuevo lote (g/t)", f"=SUM(E{l_ini}:E{l_fin})/B5", negrita, "0.0000"),
    ]
    for i, (et, val, fuente, fmt) in enumerate(resumen, start=3):
        ws.cell(i, 1, et).font = negrita
        c = ws.cell(i, 2, val)
        c.font = fuente
        c.alignment = Alignment(horizontal="right")
        if fmt:
            c.number_format = fmt
        ws.cell(i, 1).border = c.border = borde
    ws["B7"].fill = PatternFill("solid", fgColor="FFFF00")

    ws["A9"] = "Rangos de lotes"
    ws["A9"].font = negrita
    ws["A10"], ws["B10"], ws["C10"], ws["D10"] = "Rango", "Lotes ingresados", f"Peso del rango{u}", "Cantidad de lotes"
    for k, gr in enumerate(r.grupos):
        fila = g_ini + k
        ws.cell(fila, 1, gr.nombre).font = normal
        ws.cell(fila, 2, gr.texto).font = azul
        ws.cell(fila, 3, gr.peso).font = azul
        ws.cell(fila, 4, f"=COUNTIF($A${l_ini}:$A${l_fin},A{fila})").font = normal
        ws.cell(fila, 3).number_format = "#,##0.00"

    ws.cell(l_cab - 1, 1, "Lotes anteriores utilizados").font = negrita
    ws.cell(l_cab - 1, 4, "Fórmula: tenor = Σ[(peso del rango / n del rango) × tenor_i] / peso total").font = \
        Font(name=f, italic=True, color="595959")
    ws.cell(l_cab, 1, "Rango"); ws.cell(l_cab, 2, "Lote"); ws.cell(l_cab, 3, "Tenor Au (g/t)")
    ws.cell(l_cab, 4, "Peso por lote"); ws.cell(l_cab, 5, "Aporte (peso por lote × tenor)")

    fila = l_ini
    for gr in r.grupos:
        for lote, ten in zip(gr.lotes, gr.tenores):
            ws.cell(fila, 1, gr.nombre).font = normal
            ws.cell(fila, 2, lote).font = azul
            ws.cell(fila, 3, ten).font = azul
            ws.cell(fila, 4, f"=INDEX($C${g_ini}:$C${g_fin},MATCH(A{fila},$A${g_ini}:$A${g_fin},0))"
                             f"/COUNTIF($A${l_ini}:$A${l_fin},A{fila})").font = normal
            ws.cell(fila, 5, f"=D{fila}*C{fila}").font = normal
            ws.cell(fila, 3).number_format = "0.0000"
            ws.cell(fila, 4).number_format = ws.cell(fila, 5).number_format = "#,##0.0000"
            fila += 1

    # Encabezados y bordes de ambas tablas
    for fila_cab, cols in ((10, 4), (l_cab, 5)):
        for col in range(1, cols + 1):
            c = ws.cell(fila_cab, col)
            c.font, c.fill, c.border = blanca, cab, borde
            c.alignment = Alignment(horizontal="center", wrap_text=True)
    for fila_i in list(range(g_ini, g_fin + 1)):
        for col in range(1, 5):
            ws.cell(fila_i, col).border = borde
    for fila_i in range(l_ini, l_fin + 1):
        for col in range(1, 6):
            ws.cell(fila_i, col).border = borde

    ws.cell(l_fin + 2, 1, "Notas: texto azul = datos de entrada (hoja OFICIAL o ingresados por el usuario); "
                          "negro = fórmulas.").font = Font(name=f, italic=True, size=9, color="595959")
    for col, ancho in zip("ABCDE", (34, 26, 20, 20, 30)):
        ws.column_dimensions[col].width = ancho
    ws.row_dimensions[10].height = ws.row_dimensions[l_cab].height = 32

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
